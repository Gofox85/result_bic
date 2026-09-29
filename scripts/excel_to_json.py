"""Seal a recruitment results sheet (.xlsx or .csv) into the app's src/data/results.json.

Every row becomes one encrypted record that only that candidate's email + roll number can open. The
derivation below must match src/lib/resultVault.js exactly.
"""

import argparse
import base64
import csv
import hashlib
import hmac
import json
import os
import re
import sys
from datetime import date
from pathlib import Path
from zipfile import BadZipFile

from cryptography.hazmat.primitives.ciphers.aead import AESGCM
from openpyxl import load_workbook
from openpyxl.utils.exceptions import InvalidFileException


FORMAT_VERSION = 1
DEFAULT_ITERATIONS = 150_000
PADDING_BLOCK = 256

HEADER_ALIASES = {
    "email": {"email", "emailid", "emailaddress", "mail", "mailid", "collegeemail", "collegemail"},
    "rollNo": {"rollno", "rollnumber", "roll", "registerno", "registernumber", "regno"},
    "name": {"name", "fullname", "candidate", "candidatename", "studentname"},
    "department": {"department", "dept", "branch"},
    "role": {"role", "team", "position", "domain"},
    "selected": {"selected", "status", "result"},
}
REQUIRED_FIELDS = {"email", "rollNo", "name", "selected"}

EMAIL_PATTERN = re.compile(r"[^\s@]+@[^\s@]+\.[^\s@]+")
ROLL_NUMBER_PATTERN = re.compile(r"[A-Z0-9]{4,20}")


def normalize_header(value):
    return "".join(character.lower() for character in str(value or "") if character.isalnum())


def optional_text(value):
    return "" if value is None else str(value).strip()


def required_text(value, field_name, row_number):
    text = optional_text(value)
    if not text:
        raise ValueError(f"Row {row_number}: missing {field_name}.")
    return text


def email_value(value, row_number):
    email = optional_text(value).lower()
    if not EMAIL_PATTERN.fullmatch(email):
        raise ValueError(f"Row {row_number}: {email!r} is not a valid email address.")
    return email


def roll_number_value(value, row_number):
    if isinstance(value, bool) or value is None:
        raise ValueError(f"Row {row_number}: missing roll number.")
    if isinstance(value, float) and value.is_integer():
        value = int(value)
    roll_no = re.sub(r"\s+", "", str(value)).upper()
    if not ROLL_NUMBER_PATTERN.fullmatch(roll_no):
        raise ValueError(f"Row {row_number}: roll number {roll_no!r} must be 4 to 20 letters or digits.")
    return roll_no


def selected_value(value, row_number):
    normalized = normalize_header(value)
    if normalized in {"true", "yes", "y", "1", "selected"}:
        return True
    if normalized in {"false", "no", "n", "0", "notselected", "rejected"}:
        return False
    raise ValueError(f"Row {row_number}: selected must be yes/no, true/false, 1/0, or selected/not selected.")


def read_rows(input_path):
    if input_path.suffix.lower() == ".csv":
        with input_path.open(newline="", encoding="utf-8-sig") as input_file:
            return [tuple(row) for row in csv.reader(input_file)]

    workbook = load_workbook(input_path, read_only=True, data_only=True)
    try:
        return list(workbook.active.iter_rows(values_only=True))
    finally:
        workbook.close()


def read_candidates(input_path):
    rows = iter(read_rows(input_path))
    headers = next(rows, None)
    if not headers:
        raise ValueError("The sheet is empty.")

    columns = {}
    for index, header in enumerate(headers):
        normalized = normalize_header(header)
        for field_name, aliases in HEADER_ALIASES.items():
            if normalized in aliases:
                if field_name in columns:
                    raise ValueError(f"The sheet has more than one column for {field_name}.")
                columns[field_name] = index
                break

    missing_fields = REQUIRED_FIELDS - columns.keys()
    if missing_fields:
        raise ValueError("Missing required columns: " + ", ".join(sorted(missing_fields)))

    candidates = []
    seen_emails = set()
    seen_roll_numbers = set()
    for row_number, row in enumerate(rows, start=2):
        if not any(value is not None and str(value).strip() for value in row):
            continue

        def cell(field_name):
            index = columns.get(field_name)
            return row[index] if index is not None and index < len(row) else None

        email = email_value(cell("email"), row_number)
        roll_no = roll_number_value(cell("rollNo"), row_number)
        if email in seen_emails:
            raise ValueError(f"Row {row_number}: duplicate email {email}.")
        if roll_no in seen_roll_numbers:
            raise ValueError(f"Row {row_number}: duplicate roll number {roll_no}.")
        seen_emails.add(email)
        seen_roll_numbers.add(roll_no)

        selected = selected_value(cell("selected"), row_number)
        result = {"rollNo": roll_no, "name": required_text(cell("name"), "name", row_number), "selected": selected}
        if selected:
            result["role"] = required_text(cell("role"), "role for a selected candidate", row_number)
        department = optional_text(cell("department"))
        if department:
            result["department"] = department
        candidates.append((email, roll_no, result))

    if not candidates:
        raise ValueError("The sheet contains no result rows.")
    return candidates


def derive_record_secrets(email, roll_no, salt, iterations):
    master = hashlib.pbkdf2_hmac("sha256", f"{email}\n{roll_no}".encode(), salt, iterations, 32)
    record_id = hmac.new(master, b"bic-results/id", "sha256").digest()[:16].hex()
    key = hmac.new(master, b"bic-results/key", "sha256").digest()
    return record_id, key


def seal_results(candidates, iterations):
    salt = os.urandom(16)
    records = {}
    for email, roll_no, result in candidates:
        record_id, key = derive_record_secrets(email, roll_no, salt, iterations)
        nonce = os.urandom(12)
        plaintext = json.dumps(result, ensure_ascii=False, separators=(",", ":")).encode()
        # Pad with trailing spaces (still valid JSON) so a record's length can't tell selected from not selected.
        plaintext += b" " * (-len(plaintext) % PADDING_BLOCK)
        records[record_id] = base64.b64encode(nonce + AESGCM(key).encrypt(nonce, plaintext, record_id.encode())).decode()

    return {
        "format": FORMAT_VERSION,
        "published": date.today().isoformat(),
        "iterations": iterations,
        "salt": base64.b64encode(salt).decode(),
        # Sorted by id, so the file doesn't keep the sheet's order (alphabetical, selected-first, ...).
        "records": dict(sorted(records.items())),
    }


def main():
    parser = argparse.ArgumentParser(description="Seal a recruitment results sheet into the app's JSON format.")
    parser.add_argument("input_file", type=Path, help="Input .xlsx workbook (first worksheet) or .csv file")
    parser.add_argument("json_file", type=Path, help="Output JSON file, normally src/data/results.json")
    parser.add_argument("--iterations", type=int, default=DEFAULT_ITERATIONS, help=argparse.SUPPRESS)
    arguments = parser.parse_args()

    try:
        candidates = read_candidates(arguments.input_file)
        vault = seal_results(candidates, arguments.iterations)
        arguments.json_file.parent.mkdir(parents=True, exist_ok=True)
        with arguments.json_file.open("w", encoding="utf-8") as output_file:
            json.dump(vault, output_file, indent=2)
            output_file.write("\n")
    except (OSError, ValueError, InvalidFileException, BadZipFile) as error:
        print(f"Error: {error}", file=sys.stderr)
        return 1

    selected = sum(1 for _, _, result in candidates if result["selected"])
    print(
        f"Sealed {len(candidates)} result(s) ({selected} selected, {len(candidates) - selected} not selected) "
        f"into {arguments.json_file}."
    )
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
